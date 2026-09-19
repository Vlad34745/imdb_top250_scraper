import logging
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

COLUMN_NAMES = ["Rank", "Title", "Year", "Rating", "IMDb ID"]


def save_to_excel(movies: list, filepath: str) -> None:
    """Writes the movie list to a styled .xlsx file at the given filepath.

    Args:
        movies: List of movie dicts, each with rank, title, year, rating,
            and movie_id keys (as returned by core.parser.parse_movies).
        filepath: Destination path for the .xlsx file. Parent directories
            are created automatically if they don't exist.

    Raises:
        ValueError: If ``movies`` is empty.
    """
    if not movies:
        raise ValueError("Cannot export an empty movie list")

    df = pd.DataFrame(movies)
    df.columns = COLUMN_NAMES

    output_dir = os.path.dirname(filepath)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    df.to_excel(filepath, index=False, sheet_name="IMDb Top 250")

    wb = load_workbook(filepath)
    ws = wb.active

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(filepath)
    logger.info(f"Saved {len(movies)} movies to {filepath}")