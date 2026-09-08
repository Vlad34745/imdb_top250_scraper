import pytest
from openpyxl import load_workbook

from core.exporter import save_to_excel

SAMPLE_MOVIES = [
    {"rank": 1, "title": "The Shawshank Redemption", "year": 1994, "rating": 9.3, "movie_id": "tt0111161"},
    {"rank": 2, "title": "The Godfather", "year": 1972, "rating": 9.2, "movie_id": "tt0068646"},
]


def test_creates_file_with_correct_headers(tmp_path):
    filepath = tmp_path / "output.xlsx"
    save_to_excel(SAMPLE_MOVIES, str(filepath))

    assert filepath.exists()
    wb = load_workbook(filepath)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == ["Rank", "Title", "Year", "Rating", "IMDb ID"]


def test_writes_correct_row_count(tmp_path):
    filepath = tmp_path / "output.xlsx"
    save_to_excel(SAMPLE_MOVIES, str(filepath))

    wb = load_workbook(filepath)
    ws = wb.active
    # header row + 2 data rows
    assert ws.max_row == 3


def test_header_row_is_styled(tmp_path):
    filepath = tmp_path / "output.xlsx"
    save_to_excel(SAMPLE_MOVIES, str(filepath))

    wb = load_workbook(filepath)
    ws = wb.active
    assert ws["A1"].font.bold is True


def test_creates_output_directory_if_missing(tmp_path):
    filepath = tmp_path / "nested" / "dir" / "output.xlsx"
    save_to_excel(SAMPLE_MOVIES, str(filepath))
    assert filepath.exists()


def test_raises_on_empty_movie_list(tmp_path):
    filepath = tmp_path / "output.xlsx"
    with pytest.raises(ValueError):
        save_to_excel([], str(filepath))
